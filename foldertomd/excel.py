# Copyright (C) 2026 Yassin Abd El Aal
# SPDX-License-Identifier: AGPL-3.0-or-later
"""Streaming stored-value spreadsheet exporter. No evaluation and no formula/VBA output.
OOXML formulas are checked only for presence; their text is never read or logged.
Sparse cells and shared strings are spooled to SQLite to bound Python memory."""
from pathlib import PurePosixPath
import datetime, html, io, json, re, sqlite3, zipfile, posixpath
from defusedxml.ElementTree import iterparse, fromstring
from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries
from openpyxl.utils.datetime import from_excel, WINDOWS_EPOCH, MAC_EPOCH
from .render import table, text, cell

N = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'

def display(value, fmt, ctx, where, epoch=WINDOWS_EPOCH):
    if isinstance(value, bool): return 'TRUE' if value else 'FALSE'
    if value is None: return ''
    if isinstance(value, str): return value
    if fmt in ('General', '@', ''): return str(value)
    if is_date_format(fmt):
        try:
            dt = from_excel(value, epoch=epoch) if isinstance(value, (int, float)) else value
            low = fmt.lower()
            if isinstance(dt, datetime.time): return dt.strftime('%H:%M:%S' if 's' in low else '%H:%M')
            if '[h]' in low:
                return str(datetime.timedelta(days=float(value)))
            datefmt = '%Y-%m-%d'
            if low in ('mm-dd-yy', 'mm/dd/yy'): datefmt = '%m/%d/%y'
            elif low in ('dd/mm/yyyy', 'd/m/yyyy'): datefmt = '%d/%m/%Y'
            elif low in ('d-mmm-yy',): datefmt = '%d-%b-%y'
            elif low in ('d-mmm',): datefmt = '%d-%b'
            if 'h' in low: datefmt += ' %H:%M:%S' if 's' in low else ' %H:%M'
            return dt.strftime(datefmt)
        except (ValueError, OverflowError, TypeError):
            ctx.warn(f'{where}: date format could not be represented; underlying value retained.'); return str(value)
    # Recognize common numeric display patterns; never pretend complex formats are exact.
    first = fmt.split(';')[0]
    clean = re.sub(r'\[[^\]]*\]', '', first)
    literals = ''.join(re.findall(r'"([^"]*)"', clean))
    clean = re.sub(r'"[^"]*"', '', clean).replace('\\', '')
    if re.fullmatch(r'0+', clean): return f'{int(value):0{len(clean)}d}' if float(value).is_integer() else str(value)
    if re.fullmatch(r'[$€£¥]?\s*[#,0]+(?:\.[0#]+)?%?', clean.strip()):
        dec = len(clean.split('.')[1].rstrip('%')) if '.' in clean else 0
        val = float(value) * (100 if '%' in clean else 1)
        rendered = format(val, (',' if ',' in clean else '') + f'.{dec}f')
        currency = next((c for c in '$€£¥' if c in clean), '')
        return currency + literals + rendered + ('%' if '%' in clean else '')
    ctx.warn(f'{where}: custom number format cannot be faithfully represented; underlying value retained.')
    return str(value)

def blocks(db, sheet, merges, ctx):
    occupied = {r[0] for r in db.execute('select distinct c from cells where s=? order by c', (sheet,))}
    for a,b,c,d in merges:
        if c-a<32 and d-b<200:
            occupied.update(range(a,c+1))
            # Materialize only small merged regions so HTML row/column spans stay valid.
            for row in range(b,d+1):
                db.execute('insert or ignore into cells values(?,?,?,?)',(sheet,row,a,''))
    cols = sorted(occupied)
    if not cols:
        ctx.write('(Empty worksheet)\n\n'); return
    groups = []
    for c in cols:
        if not groups or c > groups[-1][-1]+1 or len(groups[-1]) >= 32: groups.append([])
        groups[-1].append(c)
    for columns in groups:
        lo, hi = columns[0], columns[-1]
        cursor = db.execute('select r,c,v from cells where s=? and c between ? and ? order by r,c', (sheet, lo, hi))
        buffer, current, values = [], None, {}
        def flush():
            if not buffer: return
            a, b = buffer[0][0], buffer[-1][0]
            ctx.write(f'### {get_column_letter(lo)}{a}:{get_column_letter(hi)}{b}\n\n')
            relevant = [m for m in merges if m[0]>=lo and m[2]<=hi and m[1]>=a and m[3]<=b]
            if relevant:
                ctx.write('<table>\n<tr><th>Row</th>'+''.join(f'<th>{get_column_letter(c)}</th>' for c in range(lo,hi+1))+'</tr>\n')
                for row, vals in buffer:
                    ctx.write(f'<tr><th>{row}</th>')
                    for c in range(lo, hi+1):
                        m = next((m for m in relevant if m[0]<=c<=m[2] and m[1]<=row<=m[3]), None)
                        if m and (c,row)!=(m[0],m[1]): continue
                        attrs = f' colspan="{m[2]-m[0]+1}" rowspan="{m[3]-m[1]+1}"' if m else ''
                        ctx.write(f'<td{attrs}>'+html.escape(vals.get(c, '')).replace('\n','<br>')+'</td>')
                    ctx.write('</tr>\n')
                ctx.write('</table>\n\n')
            else:
                ctx.write(table([['Row']+[get_column_letter(c) for c in range(lo,hi+1)]] + [[str(row)]+[vals.get(c,'') for c in range(lo,hi+1)] for row,vals in buffer]))
            buffer.clear()
        for row,c,value in cursor:
            if row != current:
                if current is not None: buffer.append((current, values))
                if buffer and (row > current+1 or len(buffer)>=200): flush()
                current, values = row, {}
            values[c] = value
        if current is not None: buffer.append((current, values))
        flush()
    if merges:
        ctx.write('Merged ranges: '+', '.join(f'{get_column_letter(a)}{b}:{get_column_letter(c)}{d}' for a,b,c,d in merges)+'.\n\n')
        # Ranges spanning separate blocks retain exact coordinates instead of creating huge blank grids.

def office_images(z, prefix, ctx):
    from PIL import Image
    for name in sorted(z.namelist()):
        if name.startswith(prefix+'/media/') and not name.endswith('/'):
            try:
                with Image.open(io.BytesIO(z.read(name))) as img:
                    ctx.image(img, 'Embedded worksheet image')
            except Exception:
                ctx.warn('An embedded vector or unsupported image could not be rendered.')
    if any(n.startswith(prefix+'/charts/') for n in z.namelist()):
        ctx.warn('Chart visuals are not rendered; worksheet stored values were exported. Chart expressions are never exported.')
    if any(n.startswith(prefix+'/embeddings/') for n in z.namelist()):
        ctx.warn('Embedded OLE objects are not extracted or executed.')

def extract_xlsx(ctx):
    db = sqlite3.connect(ctx.work/'values.sqlite')
    db.execute('create table strings (i integer primary key,v text)')
    db.execute('create table cells(s integer,r integer,c integer,v text,primary key(s,r,c)) without rowid')
    with zipfile.ZipFile(ctx.source) as z:
        names = set(z.namelist())
        if 'xl/sharedStrings.xml' in names:
            index = 0
            with z.open('xl/sharedStrings.xml') as f:
                for _, e in iterparse(f, events=('end',)):
                    if e.tag == N+'si':
                        db.execute('insert into strings values(?,?)',(index,''.join(t.text or '' for t in e.iter(N+'t')))); index+=1; e.clear()
        formats = dict(BUILTIN_FORMATS); styles = ['General']
        if 'xl/styles.xml' in names:
            root = fromstring(z.read('xl/styles.xml'))
            for f in root.findall(f'{N}numFmts/{N}numFmt'): formats[int(f.attrib['numFmtId'])] = f.attrib['formatCode']
            styles = [formats.get(int(x.get('numFmtId','0')), 'General') for x in root.findall(f'{N}cellXfs/{N}xf')]
        wb = fromstring(z.read('xl/workbook.xml'))
        pr = wb.find(N+'workbookPr')
        epoch = MAC_EPOCH if pr is not None and pr.get('date1904') in ('1','true') else WINDOWS_EPOCH
        rels = {r.get('Id'): r.get('Target') for r in fromstring(z.read('xl/_rels/workbook.xml.rels')) if r.get('TargetMode') != 'External'}
        sheets = wb.findall(f'{N}sheets/{N}sheet'); ctx.counts['sheets'] = len(sheets)
        for si, sheet in enumerate(sheets):
            name = sheet.get('name'); state = sheet.get('state','visible')
            ctx.stage(f'Reading worksheet {si+1} of {len(sheets)}',si+1,len(sheets),'Sheet')
            ctx.write('## Worksheet: '+text(name)+(f' ({state})' if state!='visible' else '')+'\n\n')
            target = rels.get(sheet.get(R+'id'))
            if not target: ctx.warn(f'{name}: worksheet relationship unavailable'); continue
            path = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/'+target)
            merges = []
            with z.open(path) as f:
                for _,e in iterparse(f, events=('end',)):
                    if e.tag == N+'c':
                        coord = e.get('r'); row,col = coordinate_to_tuple(coord)
                        v = e.find(N+'v'); form = e.find(N+'f')
                        typ = e.get('t','n'); raw = v.text if v is not None else None
                        if form is not None and (v is None or (raw is None and typ != 'str')):
                            value = '[Saved value unavailable]'; ctx.warn(f'{name}!{coord}: saved value unavailable.')
                        elif typ == 'inlineStr': value = ''.join(t.text or '' for t in e.iter(N+'t'))
                        elif typ == 's' and raw is not None:
                            entry = db.execute('select v from strings where i=?',(int(raw),)).fetchone(); value = entry[0] if entry else '[Missing shared string]'
                        elif typ == 'b' and raw is not None: value = raw != '0'
                        elif typ in ('str','e','d'): value = raw or ''
                        elif raw is not None:
                            value = float(raw); value = int(value) if value.is_integer() else value
                        else: value = None
                        if value is not None:
                            style = int(e.get('s','0')); fmt = styles[style] if style<len(styles) else 'General'
                            val = display(value, fmt, ctx, f'{name}!{coord}', epoch)
                            db.execute('insert or replace into cells values(?,?,?,?)',(si,row,col,val))
                        e.clear()
                    elif e.tag == N+'mergeCell': merges.append(range_boundaries(e.attrib['ref'])); e.clear()
                    elif e.tag == N+'row': e.clear()
            db.commit(); blocks(db, si, merges, ctx)
        office_images(z, 'xl', ctx)
    db.close()

def extract_xls(ctx):
    import xlrd
    db = sqlite3.connect(ctx.work/'values.sqlite')
    db.execute('create table cells(s integer,r integer,c integer,v text,primary key(s,r,c)) without rowid')
    book = xlrd.open_workbook(str(ctx.source), on_demand=True, formatting_info=True, ragged_rows=True)
    ctx.counts['sheets'] = book.nsheets
    try:
        for i in range(book.nsheets):
            s = book.sheet_by_index(i); name=s.name
            ctx.stage(f'Reading worksheet {i+1} of {book.nsheets}',i+1,book.nsheets,'Sheet')
            state = {0:'',1:' (hidden)',2:' (veryHidden)'}.get(s.visibility,'')
            ctx.write('## Worksheet: '+text(name)+state+'\n\n')
            for r in range(s.nrows):
                for c in range(s.row_len(r)):
                    value = s.cell_value(r,c); typ=s.cell_type(r,c)
                    if typ in (xlrd.XL_CELL_EMPTY,xlrd.XL_CELL_BLANK): continue
                    if typ==xlrd.XL_CELL_BOOLEAN: value=bool(value)
                    elif typ==xlrd.XL_CELL_ERROR: value=xlrd.error_text_from_code.get(value,'#ERROR!')
                    fmt = book.format_map[book.xf_list[s.cell_xf_index(r,c)].format_key].format_str
                    val=display(value,fmt,ctx,f'{name}!{get_column_letter(c+1)}{r+1}',MAC_EPOCH if book.datemode else WINDOWS_EPOCH)
                    db.execute('insert into cells values(?,?,?,?)',(i,r+1,c+1,val))
            db.commit()
            merges=[(cl+1,rl+1,ch,rh) for rl,rh,cl,ch in s.merged_cells]
            blocks(db,i,merges,ctx); book.unload_sheet(i)
        from .legacy import extract_xls_images
        extract_xls_images(ctx)
        ctx.warn('Legacy XLS charts and complex drawing shapes are not rendered; cell values and embedded raster images are retained.')
    finally: book.release_resources(); db.close()
