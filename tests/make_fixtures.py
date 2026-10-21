# Copyright (C) 2026 Yassin Abd El Aal
# SPDX-License-Identifier: AGPL-3.0-or-later
"""Synthetic test documents only; no user documents or external services."""
from pathlib import Path
import datetime, io, json, shutil, struct, sys, zipfile
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import pymupdf as fitz
from PIL import Image,ImageDraw,ImageFont
import arabic_reshaper
from bidi.algorithm import get_display
import xlsxwriter, xlwt
from docx import Document
from docx.shared import Inches
from pptx import Presentation
from pptx.util import Inches as PInches

def rewrite_zip(path,transform,extras=None):
    with zipfile.ZipFile(path) as z: contents={n:z.read(n) for n in z.namelist()}
    for name,data in list(contents.items()): contents[name]=transform(name,data)
    contents.update(extras or {})
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for name,data in contents.items(): z.writestr(name,data)

HIDDEN_SENTINEL='HIDDEN_GARBLED_OCR_LAYER_9147'
HIDDEN_ARABIC='ملاعلاب ابحرم'

def hidden_ocr_pdf(path,imagebytes,full_page=True):
    """A genuine raster scan plus invisible searchable text, never rasterized text."""
    with fitz.open() as pdf:
        page=pdf.new_page(width=600,height=500)
        rect=fitz.Rect(0,0,600,500) if full_page else fitz.Rect(60,90,540,360)
        page.insert_image(rect,stream=imagebytes,keep_proportion=False)
        page.insert_font(fontname='fixtureArabic',fontfile='C:/Windows/Fonts/arial.ttf')
        # Place hidden words over the real Arabic scan so masking would erase it.
        y=rect.y0+rect.height*0.69
        page.insert_text((rect.x0+rect.width*.055,y),HIDDEN_ARABIC,
                         fontname='fixtureArabic',fontsize=rect.width*.036,render_mode=3)
        page.insert_text((rect.x0+20,rect.y0+35),HIDDEN_SENTINEL,fontsize=12,render_mode=3)
        if not full_page:page.insert_text((25,35),'VISIBLE HEADER 6721',fontsize=17)
        pdf.save(path)

def broken_text_pdf(path,mixed=False):
    """Visible vector glyphs with deliberately broken Unicode extraction, no scan."""
    with fitz.open() as pdf:
        page=pdf.new_page(width=600,height=500)
        font=fitz.Font(fontfile='C:/Windows/Fonts/arial.ttf')
        xref=page.insert_font(fontname='broken',fontbuffer=font.buffer)
        phrase='RECOVERED FONT TEXT 4826'
        page.insert_text((35,240),phrase,fontname='broken',fontsize=25)
        cmap=int(pdf.xref_get_key(xref,'ToUnicode')[1].split()[0])
        entries='\n'.join(f'<{font.has_glyph(ord(c)):04x}> <0001>' for c in sorted(set(phrase)) if c!=' ')
        data='/CIDInit /ProcSet findresource begin 12 dict begin begincmap\n/CMapType 2 def\n1 begincodespacerange\n<0000> <ffff>\nendcodespacerange\n'
        data+=f'{len(set(phrase)-{" "})} beginbfchar\n'+entries+'\nendbfchar\nendcmap CMapName currentdict /CMap defineresource pop end end'
        pdf.update_stream(cmap,data.encode('ascii'))
        if mixed:
            # Keep the page-wide broken-character share below 5%, but this block above 20%.
            for i in range(12):page.insert_text((25,30+i*12),'Readable native paragraph for block fallback verification. '+str(i),fontsize=9)
        page.add_text_annot((560,450),'Annotation retained 7381')
        pdf.save(path)

def rotated_image_pdf(path):
    with fitz.open() as pdf:
        page=pdf.new_page(width=400,height=800)
        image=Image.new('RGB',(1400,500),'white')
        draw=ImageDraw.Draw(image);font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',64)
        for i,line in enumerate(['ROTATED IMAGE TEXT 7392','Correct clipping preserves this image','Readable rotated page content']):
            draw.text((35,40+i*130),line,font=font,fill='black')
        # After page rotation this pre-rotated raster reads normally; old unrotated
        # clip y=500..780 was entirely outside the rotated page's height of 400.
        stream=io.BytesIO();image.rotate(90,expand=True).save(stream,'PNG')
        page.insert_image(fitz.Rect(50,500,150,780),stream=stream.getvalue())
        page.set_rotation(90)
        pdf.save(path)

def make(destination,large=False):
    dest=Path(destination); dest.mkdir(parents=True,exist_ok=True)
    (dest/'nested'/'empty').mkdir(parents=True,exist_ok=True)
    (dest/'genuine_md').mkdir(exist_ok=True)
    (dest/'genuine_md'/'keep.txt').write_text('Legitimate folder retained',encoding='utf-8')
    (dest/'nested'/'ملاحظات.txt').write_text('نص عربي للاختبار\nEnglish text | preserved\n<not active HTML>',encoding='utf-8')
    (dest/'notes.md').write_bytes(b'# Original Markdown\r\n\r\nExact bytes.\r\n')
    (dest/'table.csv').write_text('Name,Value\n"A|B","line 1\nline 2"\nzero,0\n',encoding='utf-8')
    (dest/'table.tsv').write_text('Column\tNumber\nالكمية\t25\n',encoding='utf-8')
    (dest/'page.html').write_text('<h1>Safe HTML</h1><p>Visible text</p><script>SECRET_SCRIPT</script><img src="https://invalid.example/image.png"><a href="javascript:alert(1)">Link label</a>',encoding='utf-8')
    shutil.copyfile(dest/'page.html',dest/'page.htm')
    (dest/'unsupported.bin').write_bytes(b'unsupported fixture')
    (dest/'~$locked.docx').write_bytes(b'office lock')
    (dest/'broken.docx').write_bytes(b'not a zip')
    scan=Image.new('RGB',(1600,900),'white'); draw=ImageDraw.Draw(scan)
    font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',54)
    lines=['ENGLISH SCANNED BODY','Invoice number 4826','Quantity 25   Price 120.50','OFFLINE DOCUMENT CONVERSION']
    for i,line in enumerate(lines): draw.text((90,70+i*110),line,font=font,fill='black')
    # Fixture rasterization needs shaping when the Windows Pillow wheel lacks libraqm.
    # This is fixture-only; extraction never reverses/reorders Arabic OCR output.
    draw.text((90,570),get_display(arabic_reshaper.reshape('مرحبا بالعالم')),font=font,fill='black')
    draw.text((90,680),get_display(arabic_reshaper.reshape('فاتورة رقم ١٢٣')),font=font,fill='black')
    scan.save(dest/'scan.png',dpi=(300,300))
    for ext,fmt in [('jpg','JPEG'),('jpeg','JPEG'),('bmp','BMP'),('webp','WEBP')]: scan.save(dest/('scan.'+ext),format=fmt)
    second=Image.new('RGB',(1600,900),'white');ImageDraw.Draw(second).text((80,180),'SECOND TIFF PAGE 7392',font=font,fill='black')
    scan.save(dest/'multipage.tiff',save_all=True,append_images=[second],compression='tiff_deflate')
    scan.save(dest/'single.tif',compression='tiff_deflate')
    imagebytes=(dest/'scan.png').read_bytes()
    hidden_ocr_pdf(dest/'hidden-ocr.pdf',imagebytes)
    broken_text_pdf(dest/'control-text.pdf')
    broken_text_pdf(dest/'control-block.pdf',mixed=True)
    rotated_image_pdf(dest/'rotated-image.pdf')
    pdf=fitz.open();p=pdf.new_page();p.insert_text((50,60),'Native heading 7612',fontsize=20);p.insert_text((50,110),'Native paragraph with useful contents.',fontsize=12)
    for y in [170,200,230,260]:p.draw_line((50,y),(350,y))
    for x in [50,200,350]:p.draw_line((x,170),(x,260))
    for y,items in zip([190,220,250],[['Item','Value'],['Alpha','125'],['Beta','250']]):
        for x,word in zip([60,210],items):p.insert_text((x,y),word)
    pdf.save(dest/'same.pdf');pdf.close()
    pdf=fitz.open();p=pdf.new_page(width=600,height=400);p.insert_image(p.rect,stream=imagebytes);pdf.save(dest/'scanned.pdf');pdf.close()
    pdf=fitz.open();p=pdf.new_page(width=600,height=450);p.insert_text((25,30),'NATIVE HEADER 9031',fontsize=17);p.insert_image(fitz.Rect(20,65,580,410),stream=imagebytes)
    p=pdf.new_page(width=450,height=600);rot=io.BytesIO();scan.rotate(90,expand=True).save(rot,'PNG');p.insert_image(p.rect,stream=rot.getvalue())
    pdf.save(dest/'mixed.pdf');pdf.close()
    pdf=fitz.open();p=pdf.new_page();p.insert_text((30,40),'Protected test');pdf.save(dest/'protected.pdf',encryption=fitz.PDF_ENCRYPT_AES_256,user_pw='fixture-password',owner_pw='fixture-owner');pdf.close()
    doc=Document();doc.add_heading('Document heading 6394',level=1);doc.add_paragraph('First list item',style='List Bullet');doc.add_paragraph('Second list item',style='List Number')
    tab=doc.add_table(rows=2,cols=2);tab.cell(0,0).text='Item';tab.cell(0,1).text='Amount';tab.cell(1,0).text='Pipe | value';tab.cell(1,1).text='12.50'
    doc.add_picture(str(dest/'scan.png'),width=Inches(5));doc.sections[0].header.paragraphs[0].text='Header reference 4139';doc.sections[0].footer.paragraphs[0].text='Footer reference 8217'
    doc.save(dest/'same.docx')
    shutil.copyfile(dest/'same.docx',dest/'macro.docm')
    rewrite_zip(dest/'macro.docm',lambda n,d:d.replace(b'application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml',b'application/vnd.ms-word.document.macroEnabled.main+xml') if n=='[Content_Types].xml' else d,{'word/vbaProject.bin':b'VBA_SENTINEL_MUST_NOT_APPEAR'})
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[5]);slide.shapes.title.text='Slide heading 5908'
    box=slide.shapes.add_textbox(PInches(.6),PInches(1),PInches(5),PInches(.7));box.text='Slide body content'
    table=slide.shapes.add_table(2,2,PInches(.5),PInches(1.7),PInches(5),PInches(1)).table
    table.cell(0,0).text='Label';table.cell(0,1).text='Value';table.cell(1,0).text='Test';table.cell(1,1).text='876'
    slide.shapes.add_picture(str(dest/'scan.png'),PInches(.5),PInches(3),width=PInches(5))
    slide.notes_slide.notes_text_frame.text='Speaker notes 2178'
    prs.save(dest/'slides.pptx');shutil.copyfile(dest/'slides.pptx',dest/'macro.pptm')
    rewrite_zip(dest/'macro.pptm',lambda n,d:d.replace(b'application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml',b'application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml') if n=='[Content_Types].xml' else d,{'ppt/vbaProject.bin':b'VBA_SENTINEL_MUST_NOT_APPEAR'})
    book=xlsxwriter.Workbook(dest/'same.xlsx');s=book.add_worksheet('Main'); hidden=book.add_worksheet('Hidden');hidden.hide();very=book.add_worksheet('Very hidden');very.very_hidden()
    s.write_row('A1',['Kind','Saved value']);s.write('A2','Leading zeros');s.write_string('B2','00123');s.write('A3','Actual zero');s.write_number('B3',0)
    s.write('A4','Currency');s.write_number('B4',1234.5,book.add_format({'num_format':'$#,##0.00'}))
    s.write('A5','Percentage');s.write_number('B5',.125,book.add_format({'num_format':'0.0%'}))
    s.write('A6','Date');s.write_datetime('B6',datetime.datetime(2026,9,22),book.add_format({'num_format':'yyyy-mm-dd'}))
    s.write('A7','Saved numeric result');s.write_formula('B7','=SUM(40,2)',None,42)
    s.write('A8','Missing saved result');s.write_formula('B8','=SUM(500,600)',None,0)
    s.write('A9','Literal equals');s.write_string('B9','=literal text')
    s.write('A10','Saved error');s.write_formula('B10','=1/0',None,'#DIV/0!')
    s.merge_range('D2:F2','Merged title');s.write('D3','Sparse block');s.write('F3',19.25)
    s.write('A14','Blank separator above');s.write('B14',7,book.add_format({'num_format':'00000'}))
    s.write('Z10000','Far sparse value');s.set_row(1048575,20,book.add_format({'bg_color':'#FFFF00'}))
    s.insert_image('D6',str(dest/'scan.png'),{'x_scale':.1,'y_scale':.1})
    hidden.write('A1','Hidden saved value');very.write('A1','Very hidden saved value');book.close()
    def missing(n,d):
        if n=='xl/worksheets/sheet1.xml':
            d=d.replace(b'<c r="B8"><f>SUM(500,600)</f><v>0</v></c>',b'<c r="B8"><f>SUM(500,600)</f></c>')
        return d
    rewrite_zip(dest/'same.xlsx',missing)
    shutil.copyfile(dest/'same.xlsx',dest/'macro.xlsm')
    rewrite_zip(dest/'macro.xlsm',lambda n,d:d.replace(b'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml',b'application/vnd.ms-excel.sheet.macroEnabled.main+xml') if n=='[Content_Types].xml' else d,{'xl/vbaProject.bin':b'VBA_SENTINEL_MUST_NOT_APPEAR'})
    old=xlwt.Workbook();s=old.add_sheet('Legacy values');s.write(0,0,'Legacy saved values');s.write(1,0,'00123');s.write(2,0,0);s.write(3,0,123.5,xlwt.easyxf(num_format_str='$#,##0.00'));s.write(4,0,xlwt.Formula('40+2'));h=old.add_sheet('Hidden legacy');h.visibility=1;h.write(0,0,'Hidden legacy text');old.save(str(dest/'legacy.xls'))
    data=(dest/'legacy.xls').read_bytes(); needle=struct.pack('<HHHH',6,29,4,0)
    offset=data.find(needle)
    if offset!=-1: data=data[:offset+10]+struct.pack('<d',42)+data[offset+18:];(dest/'legacy.xls').write_bytes(data)
    (dest/'legacy.rtf').write_text('{\\rtf1\\ansi Legacy RTF content 7142\\par Second paragraph}',encoding='ascii')
    if large:
        big=xlsxwriter.Workbook(dest/'large.xlsx',{'constant_memory':True});s=big.add_worksheet('Large')
        for r in range(50000):s.write_row(r,0,[r,'Row '+str(r),r/100])
        big.close()
    return dest

if __name__=='__main__':
    import argparse
    p=argparse.ArgumentParser();p.add_argument('destination');p.add_argument('--large',action='store_true');args=p.parse_args();make(args.destination,args.large)
